#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
reparer_formules.py — resynchronise TOUTES les formules calculées du classeur
(Charges_Fixes.Total, Charges_Variables.Total, Synthese_Camion/Societe/Entreprise),
sur TOUS les mois déjà présents, sans toucher à la moindre donnée saisie/importée.

À usage ponctuel : corrige d'un coup toute corruption historique (ex: la plage
Total="=SUM(D:F)" au lieu de "=SUM(D:G)" trouvée sur les lignes d'août 2026) et
ajoute la garde "N/D" partout où elle manquait, sans attendre que
maj_mensuelle_flotte.py retraite chaque mois individuellement.
"""
import sys
import openpyxl
from openpyxl.utils.cell import range_boundaries

from maj_mensuelle_flotte import ajouter_lignes_table, get_table_ws


def cles_existantes(ws, table_name, cle_dedup):
    """Relit toutes les valeurs de clé de dédoublonnage déjà présentes dans une table."""
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    headers = [ws.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]
    col_index = {h: i for i, h in enumerate(headers)}
    lignes = []
    for r in range(min_row + 1, max_row + 1):
        vals = {k: ws.cell(row=r, column=min_col + col_index[k]).value for k in cle_dedup if k in col_index}
        if any(v is not None for v in vals.values()):
            lignes.append(vals)
    return lignes


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 reparer_formules.py TABLEAU_DE_BORD_FLOTTE.xlsx")
        sys.exit(1)
    chemin = sys.argv[1]
    wb = openpyxl.load_workbook(chemin)

    formule_total_cv = lambda r: f"=SUM(D{r}:G{r})"
    formule_total_cf = lambda r: f"=SUM(D{r}:G{r})"
    formules_sc = {
        "Société": lambda r: f'=IFERROR(INDEX(REF_Camions!$B:$B,MATCH(B{r},REF_Camions!$A:$A,0)),"")',
        "Charges_Fixes": lambda r: (f'=IF(COUNTIFS(Charges_Fixes!$B:$B,B{r},Charges_Fixes!$A:$A,A{r})=0,"N/D",'
                                     f"SUMIFS(Charges_Fixes!$I:$I,Charges_Fixes!$B:$B,B{r},Charges_Fixes!$A:$A,A{r})-F{r})"),
        "Charges_Variables": lambda r: (f'=IF(COUNTIFS(Charges_Variables!$B:$B,B{r},Charges_Variables!$A:$A,A{r})=0,"N/D",'
                                         f"SUMIFS(Charges_Variables!$I:$I,Charges_Variables!$B:$B,B{r},Charges_Variables!$A:$A,A{r}))"),
        "Charges_Mutualisees": lambda r: (f'=IF(COUNTIFS(Charges_Fixes!$B:$B,B{r},Charges_Fixes!$A:$A,A{r})=0,"N/D",'
                                           f"SUMIFS(Charges_Fixes!$G:$G,Charges_Fixes!$B:$B,B{r},Charges_Fixes!$A:$A,A{r}))"),
        "Charges_Totales": lambda r: f'=IF(OR(D{r}="N/D",E{r}="N/D",F{r}="N/D"),"N/D",D{r}+E{r}+F{r})',
        "KM": lambda r: (f'=IF(COUNTIFS(KM!$B:$B,B{r},KM!$A:$A,A{r},KM!$D:$D,"<>")=0,"N/D",'
                         f"SUMIFS(KM!$D:$D,KM!$B:$B,B{r},KM!$A:$A,A{r}))"),
        "Cout_au_km": lambda r: f'=IF(OR(G{r}="N/D",H{r}="N/D",H{r}=0),"N/D",G{r}/H{r})',
        "CA": lambda r: (f'=IF(COUNTIFS(CA!$B:$B,B{r},CA!$A:$A,A{r},CA!$D:$D,"<>")=0,"N/D",'
                         f"SUMIFS(CA!$D:$D,CA!$B:$B,B{r},CA!$A:$A,A{r}))"),
        "Resultat": lambda r: f'=IF(OR(J{r}="N/D",G{r}="N/D"),"N/D",J{r}-G{r})',
        "Marge_%": lambda r: f'=IF(OR(K{r}="N/D",J{r}=0),"N/D",K{r}/J{r})',
    }
    formules_ss = {
        "Charges_Fixes": lambda r: (f'=IF(COUNTIFS(Charges_Fixes!$C:$C,B{r},Charges_Fixes!$A:$A,A{r})=0,"N/D",'
                                     f"SUMIFS(Charges_Fixes!$I:$I,Charges_Fixes!$C:$C,B{r},Charges_Fixes!$A:$A,A{r})-E{r})"),
        "Charges_Variables": lambda r: (f'=IF(COUNTIFS(Charges_Variables!$C:$C,B{r},Charges_Variables!$A:$A,A{r})=0,"N/D",'
                                         f"SUMIFS(Charges_Variables!$I:$I,Charges_Variables!$C:$C,B{r},Charges_Variables!$A:$A,A{r}))"),
        "Charges_Mutualisees": lambda r: (f'=IF(COUNTIFS(Charges_Fixes!$C:$C,B{r},Charges_Fixes!$A:$A,A{r})=0,"N/D",'
                                           f"SUMIFS(Charges_Fixes!$G:$G,Charges_Fixes!$C:$C,B{r},Charges_Fixes!$A:$A,A{r}))"),
        "Charges_Totales": lambda r: f'=IF(OR(C{r}="N/D",D{r}="N/D",E{r}="N/D"),"N/D",C{r}+D{r}+E{r})',
        "CA": lambda r: (f'=IF(COUNTIFS(CA!$C:$C,B{r},CA!$A:$A,A{r},CA!$D:$D,"<>")=0,"N/D",'
                         f"SUMIFS(CA!$D:$D,CA!$C:$C,B{r},CA!$A:$A,A{r}))"),
        "Resultat": lambda r: f'=IF(OR(G{r}="N/D",F{r}="N/D"),"N/D",G{r}-F{r})',
        "Marge_%": lambda r: f'=IF(OR(H{r}="N/D",G{r}=0),"N/D",H{r}/G{r})',
    }
    formules_se = {
        "CA_Global": lambda r: f'=IF(COUNTIFS(CA!$A:$A,A{r},CA!$D:$D,"<>")=0,"N/D",SUMIFS(CA!$D:$D,CA!$A:$A,A{r}))',
        "Charges_Mutualisees": lambda r: (f'=IF(COUNTIFS(Charges_Fixes!$A:$A,A{r})=0,"N/D",'
                                           f"SUMIFS(Charges_Fixes!$G:$G,Charges_Fixes!$A:$A,A{r}))"),
        "Charges_Totales_Flotte": lambda r: (
            f'=IF(AND(COUNTIFS(Charges_Fixes!$A:$A,A{r})=0,COUNTIFS(Charges_Variables!$A:$A,A{r})=0),"N/D",'
            f"SUMIFS(Charges_Fixes!$I:$I,Charges_Fixes!$A:$A,A{r})"
            f"+SUMIFS(Charges_Variables!$I:$I,Charges_Variables!$A:$A,A{r}))"),
        "Resultat": lambda r: f'=IF(OR(B{r}="N/D",D{r}="N/D"),"N/D",B{r}-D{r})',
        "Marge_%": lambda r: f'=IF(OR(E{r}="N/D",B{r}=0),"N/D",E{r}/B{r})',
    }

    cibles = [
        ("Charges_Fixes", "T_ChargesFixes", ("Mois", "Camion_ID", "Société"), {"Total": formule_total_cf}),
        ("Charges_Variables", "T_ChargesVariables", ("Mois", "Camion_ID", "Société"), {"Total": formule_total_cv}),
        ("Synthese_Camion", "T_SyntheseCamion", ("Mois", "Camion_ID"), formules_sc),
        ("Synthese_Societe", "T_SyntheseSociete", ("Mois", "Société"), formules_ss),
        ("Synthese_Entreprise", "T_SyntheseEntreprise", ("Mois",), formules_se),
    ]

    total_repare = 0
    for feuille, table, cle_dedup, colonnes_formule in cibles:
        ws = get_table_ws(wb, table)
        lignes = cles_existantes(ws, table, cle_dedup)
        _, _, rep = ajouter_lignes_table(ws, table, lignes, cle_dedup=cle_dedup, colonnes_formule=colonnes_formule)
        print(f"{feuille} : {rep} formule(s) resynchronisée(s)")
        total_repare += rep

    wb.save(chemin)
    print(f"\nTotal : {total_repare} formule(s) corrigée(s). Classeur enregistré : {chemin}")


if __name__ == "__main__":
    main()
