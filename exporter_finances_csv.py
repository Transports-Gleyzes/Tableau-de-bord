#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
exporter_finances_csv.py
========================
Pont entre TABLEAU_DE_BORD_FLOTTE.xlsx (alimenté par maj_mensuelle_flotte.py)
et le Google Sheet du tableau de bord en ligne.

Lit les tables brutes T_CA, T_KM, T_ChargesVariables et T_ChargesFixes (jamais les
synthèses calculées : openpyxl ne voit pas le résultat des formules) et écrit
FINANCES_EXPORT.csv, une ligne par Mois / Camion / Société :

    Mois;Camion_ID;Société;CA;KM;Carburant;Peages;Salaires;Entretien;Charges_Fixes;Charges_Mutualisees

  - Charges_Fixes = loyer tracteur + loyer remorque + taxe essieu (hors charge mutualisée).
  - Une donnée absente reste une cellule VIDE (jamais 0), comme dans le classeur :
    le site l'affiche « N/D ».

Déposez le CSV dans Google Drive (idéalement en l'écrivant directement dans le dossier
synchronisé par « Google Drive pour ordinateur ») : le Google Sheet l'importe chaque
matin, ou tout de suite via le menu Tableau de bord > Importer les finances depuis Drive.

Usage :
    python3 exporter_finances_csv.py --classeur TABLEAU_DE_BORD_FLOTTE.xlsx \
        --sortie "G:/Mon Drive/TABLEAU DE BORD/FINANCES_EXPORT.csv" [--depuis 2026-01]
"""
import argparse
import csv
import os
import sys
from collections import OrderedDict
from datetime import date, datetime

import openpyxl
from openpyxl.utils.cell import range_boundaries

COLONNES = ["Mois", "Camion_ID", "Société", "CA", "KM", "Carburant", "Peages", "Salaires",
            "Entretien", "Charges_Fixes", "Charges_Mutualisees"]

# colonne du classeur -> colonne du CSV, par table
CORRESPONDANCES = {
    "T_CA": {"CA": "CA"},
    "T_KM": {"KM_Parcourus": "KM"},
    "T_ChargesVariables": {"Carburant": "Carburant", "Peages": "Peages",
                           "Salaire_Chauffeur": "Salaires", "Entretien_Hors_Contrat": "Entretien"},
    "T_ChargesFixes": {"Loyer_Tracteur": "Charges_Fixes", "Loyer_Remorque": "Charges_Fixes",
                       "Taxe_Essieu_Mensuelle": "Charges_Fixes",
                       "Charge_Mutualisee_Chauffeur": "Charges_Mutualisees"},
}


def societe_canonique(s):
    """Même règle que maj_mensuelle_flotte.societe_canonique ('LPB (except.)' -> 'LPB')."""
    if not s:
        return ""
    s = str(s).split("(")[0].strip()
    return {"GLEYZES": "Gleyzes", "LPB": "LPB"}.get(s.upper(), s)


def mois_texte(v):
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v or "").strip()
    if len(s) >= 7 and s[4] in "-/" and s[:4].isdigit():
        return f"{s[:4]}-{int(s[5:7]):02d}"
    return s


def nombre(v):
    """Renvoie un float, ou None si la cellule est vide / non numérique ('N/D', formule non calculée...)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\u00a0", "").replace(" ", "").replace("€", "").replace(",", ".")
    if not s or s.startswith("="):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def lire_table(wb, nom_table):
    for ws in wb.worksheets:
        if nom_table in ws.tables:
            min_col, min_row, max_col, max_row = range_boundaries(ws.tables[nom_table].ref)
            entetes = [ws.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]
            lignes = []
            for r in range(min_row + 1, max_row + 1):
                vals = [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
                if any(v not in (None, "") for v in vals):
                    lignes.append(dict(zip(entetes, vals)))
            return lignes
    print(f"[!] Table {nom_table} introuvable : colonnes correspondantes laissées vides.")
    return []


def construire_lignes(wb, depuis=None):
    resultat = OrderedDict()
    for nom_table, corresp in CORRESPONDANCES.items():
        for l in lire_table(wb, nom_table):
            mois, camion = mois_texte(l.get("Mois")), str(l.get("Camion_ID") or "").strip()
            if not mois or not camion or (depuis and mois < depuis):
                continue
            societe = societe_canonique(l.get("Société"))
            cle = (mois, camion, societe)
            ligne = resultat.setdefault(cle, {c: None for c in COLONNES[3:]})
            for src, dst in corresp.items():
                v = nombre(l.get(src))
                if v is not None:
                    ligne[dst] = (ligne[dst] or 0) + v
    return [dict(Mois=k[0], Camion_ID=k[1], **{"Société": k[2]}, **v) for k, v in sorted(resultat.items())]


def formater(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.2f}".rstrip("0").rstrip(".") if v != int(v) else str(int(v))
    return str(v)


def main():
    ap = argparse.ArgumentParser(description="Exporte les finances du classeur Excel vers FINANCES_EXPORT.csv")
    ap.add_argument("--classeur", default="TABLEAU_DE_BORD_FLOTTE.xlsx")
    ap.add_argument("--sortie", default="FINANCES_EXPORT.csv")
    ap.add_argument("--depuis", default=None, help="Premier mois exporté (AAAA-MM). Tous les mois si omis.")
    args = ap.parse_args()

    if not os.path.exists(args.classeur):
        sys.exit(f"Classeur introuvable : {args.classeur}")
    wb = openpyxl.load_workbook(args.classeur, read_only=False, data_only=True)
    lignes = construire_lignes(wb, args.depuis)
    if not lignes:
        sys.exit("Aucune ligne trouvée : rien n'a été écrit.")

    temporaire = args.sortie + ".tmp"
    with open(temporaire, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLONNES)
        for l in lignes:
            w.writerow([formater(l[c]) for c in COLONNES])
    os.replace(temporaire, args.sortie)  # pas de fichier à moitié écrit pris par la synchro Drive

    mois = sorted({l["Mois"] for l in lignes})
    print(f"{len(lignes)} lignes exportées ({mois[0]} → {mois[-1]}) dans {args.sortie}")


if __name__ == "__main__":
    main()
